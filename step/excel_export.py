"""Write portable, human-readable Excel workbooks with openpyxl."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _safe_value(value):
    """Convert NumPy/pandas scalar values to values supported by openpyxl."""
    if value is None:
        return None
    try:
        if value != value:  # NaN
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "item"):
        try:
            return value.item()
        except (TypeError, ValueError):
            pass
    return value


def write_workbook(output_path, sheets, preview_dir=None):
    """Write ``{sheet_name: pandas.DataFrame}`` to an XLSX file.

    ``preview_dir`` is retained for backward compatibility and is intentionally
    ignored. Missing data are written as blank cells rather than artificial
    zeros.
    """
    del preview_dir
    output = Path(output_path).resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("The output file must use the .xlsx extension.")
    if not sheets:
        raise ValueError("At least one worksheet is required.")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)

    for raw_name, table in sheets.items():
        name = str(raw_name)[:31] or "Sheet"
        worksheet = workbook.create_sheet(name)
        columns = [str(column) for column in table.columns]
        worksheet.append(columns)
        for row in table.itertuples(index=False, name=None):
            worksheet.append([_safe_value(value) for value in row])

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center")

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_index, column_name in enumerate(columns, start=1):
            observed = [len(column_name)]
            for cell in worksheet[get_column_letter(column_index)][1:]:
                observed.append(len(str(cell.value)) if cell.value is not None else 0)
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(observed) + 2, 40)

    workbook.save(output)
    return output
